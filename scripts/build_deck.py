"""Fill the mandated UniHack prototype template with this project's content.

The template ships as a fixed 15-slide deck with a white body and a navy accent.
This script writes into it rather than rebuilding it, so the mandated layout,
ordering and headings survive intact — only the empty space below each heading
is filled. Diagrams are drawn as native PowerPoint shapes, not images, so they
stay editable in the deck.

Usage:
    python scripts/build_deck.py            # -> UniHack_Prototype_Deck.pptx
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "[EXT] UniHack-Protoype Template .pptx"
OUTPUT = ROOT / "UniHack_Prototype_Deck.pptx"
SHOTS = ROOT / "docs" / "screenshots"

NAVY = RGBColor(0x00, 0x38, 0x78)
INK = RGBColor(0x1F, 0x24, 0x2B)
MUTED = RGBColor(0x5A, 0x63, 0x70)
ACCENT = RGBColor(0xC2, 0x41, 0x2B)
GREEN = RGBColor(0x1B, 0x7F, 0x4B)
PANEL = RGBColor(0xF1, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BODY_TOP = 1.62          # inches: first free line under every heading
BODY_LEFT = 0.42
BODY_WIDTH = 9.16
FONT = "Segoe UI"


# --------------------------------------------------------------------- text
def _style(run, *, size: float, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def add_text(
    slide, lines: list[tuple[str, dict]], *,
    top: float = BODY_TOP, left: float = BODY_LEFT, width: float = BODY_WIDTH,
    height: float = 3.5, align=PP_ALIGN.LEFT,
):
    """Write a block of styled paragraphs. Each line is (text, style kwargs)."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, (text, style) in enumerate(lines):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.space_after = Pt(style.pop("space_after", 6))
        if "space_before" in style:
            para.space_before = Pt(style.pop("space_before"))
        run = para.add_run()
        run.text = text
        _style(run, size=style.pop("size", 12), bold=style.pop("bold", False),
               color=style.pop("color", INK))
    return box


def bullets(items: list[str], *, size: float = 12.5, color: RGBColor = INK) -> list[tuple[str, dict]]:
    return [(f"•  {item}", {"size": size, "color": color, "space_after": 7}) for item in items]


# -------------------------------------------------------------------- shapes
def add_box(slide, text: str, left: float, top: float, width: float, height: float,
            *, fill: RGBColor = PANEL, line: RGBColor = NAVY, color: RGBColor = INK,
            size: float = 10.5, bold: bool = True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    box = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line
    box.line.width = Pt(1)
    box.shadow.inherit = False
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = Emu(45720)
    frame.margin_top = frame.margin_bottom = 0
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = text
    _style(run, size=size, bold=bold, color=color)
    return box


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, *, color: RGBColor = NAVY):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(1.5)
    return line


def add_label(slide, text: str, left: float, top: float, width: float,
              *, size: float = 9, color: RGBColor = MUTED, align=PP_ALIGN.CENTER,
              bold: bool = False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.26))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    _style(run, size=size, bold=bold, color=color)
    return box


def stat_row(slide, stats: list[tuple[str, str]], *, top: float, left: float = BODY_LEFT,
             total_width: float = BODY_WIDTH, height: float = 0.82,
             value_color: RGBColor = NAVY):
    """A row of number-over-label tiles."""
    gap = 0.14
    width = (total_width - gap * (len(stats) - 1)) / len(stats)
    for index, (value, label) in enumerate(stats):
        x = left + index * (width + gap)
        tile = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(width), Inches(height)
        )
        tile.fill.solid()
        tile.fill.fore_color.rgb = PANEL
        tile.line.color.rgb = PANEL
        tile.shadow.inherit = False
        frame = tile.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_top = frame.margin_bottom = 0
        head = frame.paragraphs[0]
        head.alignment = PP_ALIGN.CENTER
        head.space_after = Pt(1)
        _style(head.add_run(), size=17, bold=True, color=value_color)
        head.runs[0].text = value
        sub = frame.add_paragraph()
        sub.alignment = PP_ALIGN.CENTER
        run = sub.add_run()
        run.text = label
        _style(run, size=8.5, color=MUTED)


# ------------------------------------------------------------------- content
def slide_brief(slide) -> None:
    add_text(slide, [
        ("Six columns in. 252 columns out.", {"size": 19, "bold": True, "color": NAVY,
                                              "space_after": 10}),
        ("A distributor's catalogue arrives as a part number, a 35-character line of "
         "trade shorthand and three brand columns that are mostly placeholder markers. "
         "The Unilog delivery template asks for a full taxonomy path, six description "
         "variants, twenty feature bullets and sixty attribute triplets. Everything in "
         "between has to be recovered — and every recovered value has to be defensible.",
         {"size": 12.5, "space_after": 12}),
        ("The Product Intelligence Engine reads that shorthand the way a product manager "
         "does. It decodes what is deterministic, infers what the catalogue itself can "
         "support, asks a language model only for what is genuinely left, then judges its "
         "own output and routes anything it cannot defend to a human — with the evidence, "
         "source and confidence attached to every single value.",
         {"size": 12.5, "space_after": 14}),
    ], height=2.6)
    stat_row(slide, [
        ("1,000", "rows processed in 7.2s"),
        ("7,330", "fields recovered"),
        ("96%", "identity coverage"),
        ("100%", "taxonomy coverage"),
        ("83%", "auto-approved"),
    ], top=4.32)
    add_label(slide, "With the model on a 100-row slice: quality 59.6 → 72.4, all six content "
                     "columns filled, 5.9 attribute triplets per product, 411 calls, zero "
                     "rate-limit failures.",
              BODY_LEFT, 5.22, BODY_WIDTH, size=9, color=MUTED)


def slide_enrichment(slide) -> None:
    add_text(slide, [
        ("Worked example — one real row from the supplied file",
         {"size": 13, "bold": True, "color": NAVY, "space_after": 8}),
        ("IN:   3MABR-7100075678  |  \"3M 775L Stikit Film P150 - Cubitron II 50 Disc/Box\"  |  "
         "brand: -- Unbranded --  |  Part_Manuf: Jam Industrial Supply LLC",
         {"size": 10.5, "color": MUTED, "space_after": 10}),
    ], top=1.55, height=1.0)

    steps = [
        ("Decode", "Trade shorthand is parsed, not guessed: form, grit P150, attachment "
                   "system, pack quantity, selling unit — no model call."),
        ("Identify", "Part_Manuf is the account bought from, not the maker. It is published "
                     "as supplier; 3M is recovered as brand and manufacturer from the description."),
        ("Classify", "Keyword classifier settles 71% of the file; the model escalates the rest "
                     "into the 26-class taxonomy."),
        ("Infer", "Catalogue peers fill gaps — but brand-scoped attributes are refused outside "
                  "a brand-scoped peer group, so 3M's series never lands on a Diablo belt."),
        ("Write", "Six commerce descriptions and the feature bullets are generated to the "
                  "template's own length targets."),
        ("Judge", "A model reviews the finished record; anything unsupported is flagged, "
                  "de-confidenced and queued for a human."),
    ]
    top = 2.42
    for index, (title, body) in enumerate(steps):
        row, col = divmod(index, 2)
        x = BODY_LEFT + col * 4.66
        y = top + row * 0.92
        add_box(slide, title, x, y, 1.02, 0.42, fill=NAVY, line=NAVY, color=WHITE, size=10.5)
        add_label(slide, body, x + 1.12, y - 0.02, 3.42, size=8.8, color=INK, align=PP_ALIGN.LEFT)


def slide_opportunities(slide) -> None:
    add_text(slide, [
        ("Different from an \"AI enrichment\" wrapper",
         {"size": 13, "bold": True, "color": NAVY, "space_after": 6}),
        *bullets([
            "Deterministic first. The decoders cover all 1,000 rows in 7.2 seconds with no "
            "API key at all. The model is the last resort, not the pipeline.",
            "Every value carries method, source, evidence and confidence — a buyer can ask "
            "\"why does this say P150?\" and get an answer, not a vibe.",
            "It scores itself. Quality is decomposed into completeness, accuracy, consistency "
            "and richness, measured before and after, so uplift is a number and not a claim.",
        ], size=11.5),
        ("How it solves the problem statement",
         {"size": 13, "bold": True, "color": NAVY, "space_after": 6, "space_before": 8}),
        *bullets([
            "252-column delivery template filled directly — column order is read from the "
            "template file at runtime, so a revised template needs no code change.",
            "Blank beats wrong: anything below the publish floor is held back and queued, "
            "because a blank cell is recoverable and a confident wrong value is not.",
        ], size=11.5),
        ("USP", {"size": 13, "bold": True, "color": NAVY, "space_after": 6, "space_before": 8}),
        ("An enrichment engine that argues with itself. Switching the model on initially made "
         "the quality score fall — because the judge had started catching two real defects "
         "nobody had checked for. Both were fixed at the source. That is the difference between "
         "filling cells and shipping a catalogue.",
         {"size": 11.5, "color": ACCENT}),
    ], height=3.9)


def slide_features(slide) -> None:
    features = [
        ("Trade-shorthand decoder", "Form, grit, dimensions, colour, base type, pack quantity "
                                    "and selling unit parsed from a 35-character line."),
        ("Identity resolution", "Supplier account vs brand vs manufacturer, separated and "
                                "published on the right lines — 552 brands recovered."),
        ("26-class taxonomy + escalation", "Keyword classifier, then LLM escalation for what "
                                           "it cannot settle. Dept / Class / Fine path filled 100%."),
        ("Catalogue knowledge graph", "Peer consensus fills gaps and exposes outliers — scoped "
                                      "so inference never crosses a brand boundary."),
        ("Commerce content generation", "Six description variants and twenty feature bullets, "
                                        "written to the template's own length targets."),
        ("LLM judge + rule engine", "Unit sanity, magnitude plausibility, internal "
                                    "contradiction and evidence-support checks."),
        ("Human review queue", "Approve / correct / reject with notes; every decision is stored "
                               "and turned into reviewer calibration."),
        ("Evidence-first export", "252-column delivery CSV, flat catalogue CSV, full audit JSON, "
                                  "open-queue CSV."),
    ]
    top = 1.66
    for index, (title, body) in enumerate(features):
        row, col = divmod(index, 2)
        x = BODY_LEFT + col * 4.66
        y = top + row * 0.86
        add_label(slide, title, x, y, 4.4, size=11, color=NAVY, align=PP_ALIGN.LEFT, bold=True)
        add_label(slide, body, x, y + 0.24, 4.4, size=8.8, color=INK, align=PP_ALIGN.LEFT)


def slide_process_flow(slide) -> None:
    stages = [
        ("Ingest", "6 columns\nheader mapping"),
        ("Classify", "keyword →\nLLM escalation"),
        ("Decode", "shorthand,\nunits, enums"),
        ("Resolve", "golden record\nsource precedence"),
        ("Validate", "rules, ranges,\ncontradictions"),
    ]
    stages_two = [
        ("Enrich", "peers → graph\n→ LLM"),
        ("Generate", "6 descriptions\n20 features"),
        ("Judge", "LLM review\nde-confidence"),
        ("Route", "auto-approve\nor human queue"),
        ("Export", "252-column\ndelivery CSV"),
    ]
    width, height, gap = 1.62, 0.72, 0.26

    for row_index, row in enumerate((stages, stages_two)):
        y = 1.78 + row_index * 1.58
        for index, (title, sub) in enumerate(row):
            x = BODY_LEFT + index * (width + gap)
            add_box(slide, title, x, y, width, height, fill=NAVY, line=NAVY, color=WHITE, size=11)
            add_label(slide, sub, x, y + height + 0.04, width, size=8, color=MUTED)
            if index < len(row) - 1:
                add_arrow(slide, x + width, y + height / 2, x + width + gap, y + height / 2)

    # Wrap from the end of row one into the start of row two, routed through the
    # gap between the rows rather than diagonally across the labels.
    row_one_bottom = 1.78 + height
    row_two_top = 1.78 + 1.58
    midline = (row_one_bottom + row_two_top) / 2
    # Drop from the box edges, not their centres, so the wrap misses the captions.
    last_x = BODY_LEFT + 4 * (width + gap) + width - 0.16
    first_x = BODY_LEFT + 0.16
    add_arrow(slide, last_x, row_one_bottom, last_x, midline)
    add_arrow(slide, last_x, midline, first_x, midline)
    add_arrow(slide, first_x, midline, first_x, row_two_top)
    add_label(slide, "↺  reviewer corrections re-enter the record and re-score it",
              BODY_LEFT, 4.72, BODY_WIDTH, size=9, color=ACCENT)
    add_label(slide, "No model call happens until the deterministic stages have taken "
                     "everything they can.", BODY_LEFT, 5.02, BODY_WIDTH, size=9.5, color=MUTED)


def slide_architecture(slide) -> None:
    layers = [
        ("Interface", ["Streamlit app — 6 pages", "Headless CLI runner", "252-column exporter"]),
        ("Intelligence", ["Trade-shorthand decoders", "Identity resolver", "Knowledge graph + RAG",
                          "LLM extract / generate / judge"]),
        ("Contract", ["attributes.json — 26 classes", "Delivery template header", "Quality scorer"]),
        ("Foundation", ["Provider-agnostic LLM client", "SQLite store + prompt cache",
                        "Rate limiter + retry"]),
    ]
    y = 1.72
    for title, items in layers:
        add_box(slide, title, BODY_LEFT, y, 1.68, 0.74, fill=NAVY, line=NAVY, color=WHITE, size=11)
        width = (BODY_WIDTH - 1.86 - 0.12 * (len(items) - 1)) / len(items)
        for index, item in enumerate(items):
            x = BODY_LEFT + 1.86 + index * (width + 0.12)
            add_box(slide, item, x, y, width, 0.74, size=9, bold=False)
        y += 0.86

    add_label(slide, "Anthropic · Gemini · OpenAI · Groq — first provider with a key wins, and a "
                     "failure falls through to the next, so one vendor's rate limit cannot stop a run.",
              BODY_LEFT, y + 0.06, BODY_WIDTH, size=9.5, color=MUTED)


def slide_technologies(slide) -> None:
    groups = [
        ("Core", "Python 3.13 · pydantic · pandas · numpy"),
        ("Intelligence", "Gemini 3.5 Flash-Lite (bulk) · Gemini 3.6 Flash (vision) · "
                         "provider-agnostic client with Anthropic / OpenAI / Groq fallback"),
        ("Retrieval & graph", "networkx knowledge graph · TF-IDF peer retrieval (scikit-learn) · "
                              "rapidfuzz alias matching"),
        ("Documents", "PyMuPDF for datasheet tables · Pillow for image evidence · openpyxl"),
        ("Interface", "Streamlit · Plotly"),
        ("Storage", "SQLite — catalogue, runs, review decisions and the prompt cache"),
        ("Quality", "pytest — 35 tests over decoding, identity and request pacing"),
    ]
    top = 1.72
    for index, (title, body) in enumerate(groups):
        y = top + index * 0.49
        add_box(slide, title, BODY_LEFT, y, 1.72, 0.4, fill=NAVY, line=NAVY, color=WHITE, size=10)
        add_label(slide, body, BODY_LEFT + 1.9, y + 0.07, BODY_WIDTH - 1.9, size=10,
                  color=INK, align=PP_ALIGN.LEFT)


def slide_cost(slide) -> None:
    add_text(slide, [
        ("The deterministic pass is the product. The model is an optional accelerant.",
         {"size": 13, "bold": True, "color": NAVY, "space_after": 10}),
    ], top=1.6, height=0.5)
    stat_row(slide, [
        ("₹0", "full 1,000-row deterministic pass"),
        ("7.2s", "end to end, no network"),
        ("~4,000", "model calls for a full LLM pass"),
        ("15 RPM", "free-tier ceiling per model"),
    ], top=2.18)
    add_text(slide, [
        *bullets([
            "Running the whole file through the model on a free tier takes about five hours — "
            "so the engine is built so that it never has to. Deterministic decoding covers all "
            "1,000 rows; the model is pointed at what is left.",
            "Every completion is cached by (provider, model, prompt) hash, so a re-run during a "
            "demo or a review cycle costs nothing at all.",
            "A paid key changes only the throughput, not the architecture: at Gemini Flash-Lite "
            "list pricing a full 1,000-row enrichment is a low single-digit dollar job.",
            "Self-hosting cost is a Streamlit process and a SQLite file — no vector database, "
            "no queue, no GPU.",
        ], size=11),
    ], top=3.28, height=2.0)


def slide_snapshots(slide) -> None:
    shots = [
        ("01_overview.jpg", "Overview — measured uplift, not a claim"),
        ("02_catalog.jpg", "Catalog — every value, its evidence and its confidence"),
        ("04_review_queue.jpg", "Review Queue — only what the engine could not settle"),
        ("05_export.jpg", "Export — the 252-column delivery contract"),
    ]
    width, height = 4.44, 1.62
    for index, (name, caption) in enumerate(shots):
        path = SHOTS / name
        if not path.is_file():
            continue
        row, col = divmod(index, 2)
        x = BODY_LEFT + col * (width + 0.28)
        y = 1.66 + row * (height + 0.44)
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(width), Inches(height))
        add_label(slide, caption, x, y + height + 0.03, width, size=8.5, color=MUTED)


def slide_future(slide) -> None:
    add_text(slide, [
        ("What is deliberately not done yet", {"size": 13, "bold": True, "color": NAVY,
                                               "space_after": 7}),
        *bullets([
            "Distributor-side columns — SKU - MY_PART_NUMBER, List Price, Prop 65 — export "
            "blank. They are commercial data no enrichment can honestly invent.",
            "290 rows classify as General Product deterministically; LLM escalation resolves "
            "about three quarters, and the rest are genuinely outside the 26-class taxonomy.",
        ], size=11.5),
        ("Next", {"size": 13, "bold": True, "color": NAVY, "space_after": 7, "space_before": 10}),
        *bullets([
            "Manufacturer datasheet ingestion at scale — the PDF table extractor and the vision "
            "path are built and wired; they need a document source per SKU.",
            "Learned confidence calibration: the review decisions already stored are the "
            "training signal for replacing the hand-set confidence constants.",
            "Taxonomy expansion driven by what lands in General Product, so the schema grows "
            "from evidence rather than guesswork.",
            "Persisted knowledge graph across runs, so peer consensus improves as the "
            "catalogue grows instead of being rebuilt per run.",
        ], size=11.5),
    ], height=3.7)


def slide_links(slide) -> None:
    add_text(slide, [
        ("GitHub public repository", {"size": 12, "bold": True, "color": NAVY, "space_after": 2}),
        ("<paste repository URL>", {"size": 11.5, "color": MUTED, "space_after": 12}),
        ("Demo video (3 minutes)", {"size": 12, "bold": True, "color": NAVY, "space_after": 2}),
        ("<paste video URL>", {"size": 11.5, "color": MUTED, "space_after": 12}),
        ("Working prototype", {"size": 12, "bold": True, "color": NAVY, "space_after": 2}),
        ("<paste Streamlit Cloud URL>", {"size": 11.5, "color": MUTED, "space_after": 12}),
    ], top=1.9, height=2.6)


BUILDERS = {
    3: slide_brief,
    4: slide_enrichment,
    5: slide_opportunities,
    6: slide_features,
    7: slide_process_flow,
    9: slide_architecture,
    10: slide_technologies,
    11: slide_cost,
    12: slide_snapshots,
    13: slide_future,
    14: slide_links,
}


def main() -> int:
    if not TEMPLATE.is_file():
        print(f"Template not found: {TEMPLATE}", file=sys.stderr)
        return 1

    deck = Presentation(str(TEMPLATE))
    for number, builder in BUILDERS.items():
        builder(deck.slides[number - 1])
    deck.save(str(OUTPUT))
    print(f"Wrote {OUTPUT.name}  ({len(BUILDERS)} slides filled)")
    print("Still to fill by hand: slide 2 (team details), slide 14 (three URLs).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
